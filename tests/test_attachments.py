# tests/test_attachments.py
import sys
from pathlib import Path
from unittest.mock import MagicMock

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "src"))

from attachments import (
    MAX_ATTACHMENT_CHARS,
    MAX_TOTAL_ATTACHMENT_CHARS,
    MAX_UPLOAD_BYTES,
    extract_text,
    resolve_attachments,
    sanitize_filename,
    save_upload,
)


def _make_pdf(text: str) -> bytes:
    """Genera un PDF minimo de una pagina con el texto dado (para tests).
    Calcula offsets del xref en runtime para que pypdf lo parsee sin reparaciones."""
    content = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n%s\nendstream" % (len(content), content),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, obj in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += (f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
            f"startxref\n{xref}\n%%EOF").encode()
    return bytes(out)


def test_extract_text_txt_utf8(tmp_path):
    p = tmp_path / "nota.txt"
    p.write_text("Hola mundo", encoding="utf-8")
    assert extract_text(p) == "Hola mundo"


def test_extract_text_txt_latin1(tmp_path):
    p = tmp_path / "nota.txt"
    p.write_bytes("Caf\xe9".encode("latin-1"))
    assert extract_text(p) == "Caf\xe9"


def test_extract_text_pdf(tmp_path):
    p = tmp_path / "doc.pdf"
    p.write_bytes(_make_pdf("Hola PDF"))
    assert extract_text(p) == "Hola PDF"


def test_extract_text_docx(tmp_path):
    import docx
    d = docx.Document()
    d.add_paragraph("Primer parrafo")
    d.add_paragraph("Segundo parrafo")
    p = tmp_path / "nota.docx"
    d.save(str(p))
    assert extract_text(p) == "Primer parrafo\nSegundo parrafo"


def test_extract_text_extension_no_soportada(tmp_path):
    p = tmp_path / "datos.bin"
    p.write_bytes(b"x")
    with pytest.raises(ValueError):
        extract_text(p)


def test_sanitize_filename_quita_rutas():
    assert sanitize_filename("../../etc/passwd") == "passwd"


def test_sanitize_filename_quita_caracteres_peligrosos():
    assert sanitize_filename("a[b]c|d.txt") == "a_b_c_d.txt"


def test_sanitize_filename_vacio_devuelve_archivo():
    assert sanitize_filename("") == "archivo"


def test_sanitize_filename_truncado_80_con_extension():
    # 100 'a' + ".txt" -> 100-char stem truncado a 76 + ".txt" = 80 chars
    assert len(sanitize_filename("a" * 100 + ".txt")) == 80
    assert sanitize_filename("a" * 100 + ".txt").endswith(".txt")


def test_sanitize_filename_extension_mas_80_no_excede_limite():
    # Extension de 81 chars: el invariante <=80 debe mantenerse igualmente
    assert len(sanitize_filename("a" * 100 + "." + "x" * 80)) == 80


def _fake_file(name: str, content: str):
    fake = MagicMock()
    fake.filename = name
    fake.save.side_effect = lambda path: Path(path).write_text(content, encoding="utf-8")
    return fake


def test_save_upload_ok(tmp_path):
    result = save_upload(_fake_file("informe.txt", "contenido x100"), "ana", tmp_path)
    assert len(result["file_id"]) == 12
    assert result["name"] == "informe.txt"
    assert result["chars"] == len("contenido x100")
    assert result["truncated"] is False
    assert "warning" not in result
    files = list((tmp_path / "ana").iterdir())
    assert len(files) == 1
    assert files[0].name.startswith(result["file_id"] + "_")


def test_save_upload_usuario_con_ruta_no_escapa(tmp_path):
    result = save_upload(_fake_file("nota.txt", "contenido"), "../../ana", tmp_path)
    # El archivo queda dentro de tmp_path, no fuera
    assert not (tmp_path.parent / "ana").exists()
    files = list((tmp_path / "ana").iterdir())
    assert len(files) == 1
    assert files[0].name.startswith(result["file_id"] + "_")


def test_save_upload_extension_no_permitida(tmp_path):
    with pytest.raises(ValueError, match="Extension no permitida"):
        save_upload(_fake_file("datos.bin", "x"), "ana", tmp_path)


def test_save_upload_demasiado_grande_borra_archivo(tmp_path):
    big = "x" * (MAX_UPLOAD_BYTES + 1)
    with pytest.raises(ValueError, match="demasiado grande"):
        save_upload(_fake_file("grande.txt", big), "ana", tmp_path)
    user_dir = tmp_path / "ana"
    assert not user_dir.exists() or not any(user_dir.iterdir())


def test_save_upload_trunca_caracteres(tmp_path):
    content = "a" * (MAX_ATTACHMENT_CHARS + 500)
    result = save_upload(_fake_file("largo.txt", content), "ana", tmp_path)
    assert result["truncated"] is True
    assert result["chars"] == MAX_ATTACHMENT_CHARS + len("\n...[truncado]")


def test_save_upload_pdf_escaneado_warning(tmp_path):
    fake = MagicMock()
    fake.filename = "escaneado.pdf"
    # PDF valido pero sin texto extraible
    fake.save.side_effect = lambda path: Path(path).write_bytes(_make_pdf(""))
    result = save_upload(fake, "ana", tmp_path)
    assert result["chars"] == 0
    assert result["warning"]


def test_save_upload_pdf_corrupto_borra_archivo(tmp_path):
    fake = MagicMock()
    fake.filename = "corrupto.pdf"
    # pypdf lanza PdfStreamError (no ValueError) al parsear bytes invalidos
    fake.save.side_effect = lambda path: Path(path).write_bytes(b"%PDF-1.4\nno es un pdf")
    with pytest.raises(ValueError, match="No se pudo extraer texto del archivo"):
        save_upload(fake, "ana", tmp_path)
    user_dir = tmp_path / "ana"
    assert not user_dir.exists() or not any(user_dir.iterdir())


def test_resolve_attachments_reemplaza_marcador(tmp_path):
    meta = save_upload(_fake_file("informe.txt", "El contenido del informe."), "ana", tmp_path)
    marker = f"[adjunto:{meta['file_id']}|informe.txt]"
    resolved = resolve_attachments(f"Resume esto\n{marker}", "ana", tmp_path)
    assert "Resume esto" in resolved
    assert 'Contenido del archivo "informe.txt"' in resolved
    assert "El contenido del informe." in resolved
    assert marker not in resolved


def test_resolve_attachments_archivo_ausente(tmp_path):
    resolved = resolve_attachments("Hola\n[adjunto:abcdef123456|perdido.pdf]", "ana", tmp_path)
    assert "[archivo adjunto no disponible: perdido.pdf]" in resolved
    assert "Hola" in resolved


def test_resolve_attachments_pdf_escaneado_sin_texto(tmp_path):
    fake = MagicMock()
    fake.filename = "escaneado.pdf"
    fake.save.side_effect = lambda path: Path(path).write_bytes(_make_pdf(""))
    meta = save_upload(fake, "ana", tmp_path)
    marker = f"[adjunto:{meta['file_id']}|escaneado.pdf]"
    resolved = resolve_attachments(marker, "ana", tmp_path)
    assert "[sin texto extraible: escaneado.pdf]" in resolved


def test_resolve_attachments_limite_total(tmp_path):
    metas = []
    for i in range(2):
        metas.append(save_upload(
            _fake_file(f"grande{i}.txt", "x" * 30000), "ana", tmp_path))
    markers = "\n".join(
        f"[adjunto:{m['file_id']}|grande{i}.txt]" for i, m in enumerate(metas))
    resolved = resolve_attachments(f"Texto\n{markers}", "ana", tmp_path)
    # El texto limpio del usuario nunca se trunca
    assert resolved.startswith("Texto")
    assert "[aviso: adjuntos truncados" in resolved
    # Los bloques de adjuntos nunca superan la cuota; el resto del mensaje es
    # texto limpio del usuario ("Texto\n" + salto entre marcadores) y el aviso,
    # que no cuentan para la cuota, asi que la cota es exacta.
    notice = "\n[aviso: adjuntos truncados por limite de caracteres]"
    clean = "Texto\n\n"
    assert len(resolved) <= MAX_TOTAL_ATTACHMENT_CHARS + len(notice) + len(clean)
